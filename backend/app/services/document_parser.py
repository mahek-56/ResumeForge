"""
Document Parser Service
SAMATRIX RESUMEFORGE 2026

Extracts raw text safely from:
- PDF (pypdf)
- DOCX (python-docx)
- TXT (UTF-8, Latin-1 fallback)
"""

import io
import os
from typing import Tuple
from pypdf import PdfReader
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text pages from PDF bytes."""
    reader = PdfReader(io.BytesIO(file_bytes))
    extracted = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            extracted.append(page_text)
    return "\n".join(extracted).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract paragraphs and table text from DOCX bytes."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs).strip()


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decode text using UTF-8 with fallback to Latin-1."""
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore").strip()


def parse_resume_document(file_bytes: bytes, filename: str) -> Tuple[bool, str, str]:
    """
    Validates file extension, extracts content, and returns (success, extracted_text, error_message).
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".pdf", ".docx", ".txt"]:
        return False, "", f"Unsupported file extension '{ext}'. Please upload PDF, DOCX, or TXT."

    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(file_bytes)
        elif ext == ".docx":
            text = extract_text_from_docx(file_bytes)
        elif ext == ".txt":
            text = extract_text_from_txt(file_bytes)
        else:
            return False, "", f"Unsupported format: {ext}"

        if not text or len(text.strip()) < 10:
            return False, "", "Extracted resume content is empty or unreadable."

        return True, text, ""
    except Exception as e:
        return False, "", f"Failed to extract document contents: {str(e)}"
